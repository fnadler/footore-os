#!/usr/bin/env python3
"""
Extrai os campos do Pedido de Venda (.xlsx) do Footlink.

Uso: python parse_pedido.py <pedido.xlsx>

O pedido tem uma aba por modalidade (Clubes/Agentes × Mensal/Anual).
A aba preenchida é a que importa; as demais vêm em branco (template).
Layout: rótulos na coluna B, valores na coluna C, a partir da linha 4.

Saída: JSON no stdout com os campos normalizados + o nome da aba usada.
NÃO decide nada jurídico — só lê. As decisões ficam na SKILL.md.
"""
import sys, json, datetime
import openpyxl


def _norm(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    best = None
    for ws in wb.worksheets:
        # conta células preenchidas na coluna C (valores)
        filled = sum(1 for r in range(4, 30) if ws.cell(r, 3).value not in (None, "", "-"))
        # aba de pedido tem "Pedido de Venda" em C2
        is_pedido = (ws.cell(2, 3).value or "").strip().lower().startswith("pedido")
        if is_pedido and (best is None or filled > best[1]):
            best = (ws, filled)
    if best is None:
        raise SystemExit("Nenhuma aba de 'Pedido de Venda' preenchida encontrada.")
    ws = best[0]

    data = {"_aba": ws.title, "_perfil": "agente" if "agente" in ws.title.lower() else "clube"}
    for r in range(4, 30):
        label = ws.cell(r, 2).value
        value = ws.cell(r, 3).value
        if label and value not in (None, ""):
            data[str(label).strip()] = _norm(value)
    return data


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("uso: python parse_pedido.py <pedido.xlsx>")
    print(json.dumps(parse(sys.argv[1]), ensure_ascii=False, indent=2))
